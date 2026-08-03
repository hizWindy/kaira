"""Phase 7 tests — data export via the CLI and via generated API endpoints.

The load-bearing test in this file is the sensitive-field-leak regression:
every format, on both paths, is opened and searched for values that live in
columns named ``password``/``api_key``/``token``/``secret``. xlsx and docx are
zip containers, so the check reads *inside* them — a substring search on the
raw file would pass on compressed bytes that still hold the secret.
"""

from __future__ import annotations

import asyncio
import datetime
import json
import zipfile
from pathlib import Path

import pytest
from typer.testing import CliRunner

from kaira.core import export as ex
from kaira.core import security as sec

runner = CliRunner()

# Values written into sensitive columns. If any of these reaches a file, the
# strip failed — they are deliberately unlike anything else in the fixture.
LEAK_VALUES = ("LEAKPASSWORD", "LEAKAPIKEY", "LEAKTOKEN", "LEAKSECRET")
LEAK_COLUMNS = ("password", "api_key", "refresh_token", "client_secret")


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _seed_database(db_path: Path) -> str:
    """Create a sqlite database with a leaky users table and a clean posts table."""
    from sqlalchemy import Column, DateTime, Integer, MetaData, String, Table, insert
    from sqlalchemy.ext.asyncio import create_async_engine

    url = f"sqlite+aiosqlite:///{db_path.as_posix()}"
    metadata = MetaData()
    users = Table(
        "users",
        metadata,
        Column("id", Integer, primary_key=True),
        Column("uuid", String),
        Column("username", String),
        Column("email", String),
        Column("password", String),
        Column("api_key", String),
        Column("refresh_token", String),
        Column("client_secret", String),
        Column("status", String),
        Column("age", Integer),
        Column("created_at", DateTime),
    )
    posts = Table(
        "posts",
        metadata,
        Column("id", Integer, primary_key=True),
        Column("uuid", String),
        Column("title", String),
    )

    async def build() -> None:
        engine = create_async_engine(url)
        async with engine.begin() as conn:
            await conn.run_sync(metadata.create_all)
            await conn.execute(
                insert(users),
                [
                    {
                        "uuid": f"u-{i}",
                        "username": f"user{i}",
                        "email": f"u{i}@example.com",
                        "password": "LEAKPASSWORD",
                        "api_key": "LEAKAPIKEY",
                        "refresh_token": "LEAKTOKEN",
                        "client_secret": "LEAKSECRET",
                        "status": "active" if i % 2 else "inactive",
                        "age": 20 + i,
                        "created_at": datetime.datetime(2026, 8, 3, 9, 0, 0),
                    }
                    for i in range(1, 13)
                ],
            )
            await conn.execute(
                insert(posts),
                [{"uuid": f"p-{i}", "title": f"Post {i}"} for i in range(1, 5)],
            )
        await engine.dispose()

    asyncio.run(build())
    return url


def _write_config(root: Path, **overrides: object) -> None:
    data: dict[str, object] = {
        "output_dir": ".",
        "models_dir": "models",
        "repositories_dir": "repositories",
        "schemas_dir": "schemas",
        "services_dir": "services",
        "routers_dir": "routers",
        "default_tier": "full",
        "generated_models": [
            {"name": "User", "fields": [], "relations": []},
            {"name": "Post", "fields": [], "relations": []},
        ],
        "embedded_models": [],
        "db_type": "sqlite",
        "api_version": "v1",
        "auth_type": "jwt",
    }
    data.update(overrides)
    (root / ".kaira.json").write_text(json.dumps(data, indent=2), encoding="utf-8")


@pytest.fixture
def project(tmp_path, monkeypatch):
    """A Kaira project with a seeded database, cwd already moved into it."""
    monkeypatch.chdir(tmp_path)
    _write_config(tmp_path)
    url = _seed_database(tmp_path / "test.db")
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.delenv("APP_ENV", raising=False)
    tmp_path.joinpath(".gitignore").write_text("venv/\n", encoding="utf-8")
    return tmp_path


@pytest.fixture
def db_url(project) -> str:
    import os

    return os.environ["DATABASE_URL"]


def _file_parts(path: Path) -> list[bytes]:
    """Raw bytes plus, for zip-container formats, every member's bytes."""
    parts = [path.read_bytes()]
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            parts.extend(archive.read(name) for name in archive.namelist())
    return parts


def assert_no_leak(path: Path) -> None:
    """Fail if any sensitive value or column name survived into *path*."""
    parts = _file_parts(path)
    for needle in LEAK_VALUES + LEAK_COLUMNS:
        blob = needle.encode()
        assert not any(blob in part for part in parts), (
            f"{needle!r} leaked into {path.name}"
        )


def collect(iterator) -> list:
    """Drain an async iterator from a synchronous test."""

    async def run() -> list:
        return [item async for item in iterator]

    return asyncio.run(run())


# ---------------------------------------------------------------------------
# core/security.py — the one exclusion list
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name",
    [
        "password",
        "PASSWORD",
        "hashed_password",
        "password_hash",
        "reset_token",
        "access_token",
        "client_secret",
        "SECRET_KEY",
        "api_key",
        "user_api_key",
    ],
)
def test_sensitive_markers_match_real_column_names(name):
    assert sec.is_sensitive_field(name)


@pytest.mark.parametrize(
    "name",
    ["username", "email", "age", "status", "created_at", "uuid", "id", "keyword"],
)
def test_ordinary_columns_are_not_sensitive(name):
    assert not sec.is_sensitive_field(name)


def test_strip_sensitive_fields_removes_only_sensitive_keys():
    row = {"username": "a", "password": "x", "api_key": "y", "age": 3}
    assert sec.strip_sensitive_fields(row) == {"username": "a", "age": 3}


def test_strip_sensitive_fields_does_not_mutate_input():
    row = {"username": "a", "password": "x"}
    sec.strip_sensitive_fields(row)
    assert "password" in row


def test_safe_field_names_preserves_order():
    assert sec.safe_field_names(["b", "password", "a"]) == ["b", "a"]


def test_export_reexports_the_same_strip_function():
    """Two import paths, one implementation — never two lists."""
    assert ex.strip_sensitive_fields is sec.strip_sensitive_fields
    assert ex.SENSITIVE_FIELD_MARKERS is sec.SENSITIVE_FIELD_MARKERS


# ---------------------------------------------------------------------------
# core/export.py — pure helpers
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "model,expected",
    [
        ("User", "users"),
        ("UserProfile", "user_profiles"),
        ("Category", "categories"),
        ("Box", "boxes"),
        ("Dish", "dishes"),
    ],
)
def test_default_table_name_matches_the_parser(model, expected):
    from kaira.core.parser import table_name

    assert ex.default_table_name(model) == expected
    assert ex.default_table_name(model) == table_name(model)


def test_scalar_coerces_types_writers_cannot_hold():
    from decimal import Decimal
    from uuid import UUID

    assert ex._scalar(None) == ""
    assert ex._scalar(Decimal("1.5")) == 1.5
    assert ex._scalar(datetime.date(2026, 8, 3)) == "2026-08-03"
    assert ex._scalar(UUID(int=0)) == "00000000-0000-0000-0000-000000000000"
    assert ex._scalar({"a": 1}) == '{"a": 1}'
    assert "bytes" in ex._scalar(b"abc")
    assert ex._scalar(7) == 7


def test_text_elides_long_values():
    assert len(ex._text("x" * 500)) <= ex._MAX_CELL_CHARS


def test_sheet_title_obeys_excel_limits():
    assert ex._sheet_title("a" * 60) == "a" * 31
    assert "/" not in ex._sheet_title("a/b:c")
    assert ex._sheet_title("") == "Sheet"


def test_resolve_fields_defaults_to_all_minus_sensitive():
    assert ex.resolve_fields(["id", "password", "email"]) == ["id", "email"]


def test_resolve_fields_allowlist_cannot_widen_to_sensitive():
    assert ex.resolve_fields(["id", "password", "email"], ["email", "password"]) == [
        "email"
    ]


def test_resolve_fields_rejects_unknown_field_without_naming_secrets():
    with pytest.raises(ex.ExportError) as err:
        ex.resolve_fields(["id", "password"], ["nope"])
    assert "password" not in str(err.value)


def test_resolve_fields_rejects_an_all_sensitive_request():
    with pytest.raises(ex.ExportError):
        ex.resolve_fields(["id", "password"], ["password"])


def test_validate_filters_rejects_sensitive_keys():
    """An equality filter on a hash is a guessing oracle, so it is refused."""
    with pytest.raises(ex.ExportError):
        ex.validate_filters(["id", "password"], {"password": "guess"})


def test_validate_filters_rejects_unknown_keys():
    with pytest.raises(ex.ExportError):
        ex.validate_filters(["id"], {"nope": "1"})


def test_validate_filters_accepts_known_keys():
    ex.validate_filters(["id", "status"], {"status": "active"})


def test_write_rows_rejects_an_unknown_format():
    async def empty():
        return
        yield  # pragma: no cover — makes this an async generator

    with pytest.raises(ex.ExportError):
        asyncio.run(ex.write_rows(empty(), "csv", "out.csv", "T"))


# ---------------------------------------------------------------------------
# core/export.py — fetch_rows against a real database
# ---------------------------------------------------------------------------


def test_fetch_rows_strips_sensitive_columns(db_url):
    rows = collect(ex.fetch_rows("User", url=db_url, table="users"))
    assert len(rows) == 12
    for column in LEAK_COLUMNS:
        assert column not in rows[0]
    assert {"id", "uuid", "username", "email", "status", "age"} <= set(rows[0])


def test_fetch_rows_batches_without_skipping_or_repeating(db_url):
    """A batch smaller than the table must still yield every row exactly once."""
    rows = collect(ex.fetch_rows("User", url=db_url, table="users", batch_size=5))
    ids = [r["id"] for r in rows]
    assert ids == sorted(ids)
    assert len(ids) == len(set(ids)) == 12


def test_fetch_rows_honours_limit_across_batches(db_url):
    rows = collect(
        ex.fetch_rows("User", url=db_url, table="users", limit=7, batch_size=3)
    )
    assert len(rows) == 7


def test_fetch_rows_limit_zero_yields_nothing(db_url):
    assert collect(ex.fetch_rows("User", url=db_url, table="users", limit=0)) == []


def test_fetch_rows_applies_equality_filters(db_url):
    rows = collect(
        ex.fetch_rows("User", url=db_url, table="users", filters={"status": "active"})
    )
    assert rows and all(r["status"] == "active" for r in rows)


def test_fetch_rows_casts_filter_values_to_the_column_type(db_url):
    """'25' against an INTEGER column must match, not compare as text."""
    rows = collect(
        ex.fetch_rows("User", url=db_url, table="users", filters={"age": "25"})
    )
    assert len(rows) == 1
    assert rows[0]["age"] == 25


def test_fetch_rows_field_allowlist_narrows_output(db_url):
    rows = collect(
        ex.fetch_rows("User", url=db_url, table="users", fields=["username", "email"])
    )
    assert list(rows[0]) == ["username", "email"]


def test_fetch_rows_allowlist_cannot_reinstate_a_secret(db_url):
    rows = collect(
        ex.fetch_rows("User", url=db_url, table="users", fields=["username", "api_key"])
    )
    assert list(rows[0]) == ["username"]


def test_fetch_rows_rejects_a_sensitive_filter(db_url):
    with pytest.raises(ex.ExportError):
        collect(
            ex.fetch_rows("User", url=db_url, table="users", filters={"password": "x"})
        )


def test_fetch_rows_reports_a_missing_table_clearly(db_url):
    with pytest.raises(ex.ExportError) as err:
        collect(ex.fetch_rows("Ghost", url=db_url, table="ghosts"))
    assert "ghosts" in str(err.value)


def test_fetch_rows_without_a_handle_or_url_fails():
    with pytest.raises(ex.ExportError):
        collect(ex.fetch_rows("User"))


def test_fetch_rows_via_an_open_session_matches_the_url_path(db_url):
    """The generated service passes ``db=`` — it must see exactly what the CLI sees."""
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    async def run() -> list:
        engine = create_async_engine(db_url)
        maker = async_sessionmaker(engine, expire_on_commit=False)
        async with maker() as session:
            rows = [r async for r in ex.fetch_rows("User", db=session, table="users")]
        await engine.dispose()
        return rows

    via_session = asyncio.run(run())
    via_url = collect(ex.fetch_rows("User", url=db_url, table="users"))
    assert via_session == via_url


# ---------------------------------------------------------------------------
# core/export.py — writers, and the leak regression across all three formats
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "docx"])
def test_writer_produces_a_file_with_no_leaked_fields(fmt, db_url, tmp_path):
    out = tmp_path / f"users.{fmt}"
    count = asyncio.run(
        ex.write_rows(
            ex.fetch_rows("User", url=db_url, table="users"),
            fmt,
            str(out),
            "User export",
        )
    )
    assert count == 12
    assert out.stat().st_size > 0
    assert_no_leak(out)


@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "docx"])
def test_writer_handles_an_empty_result_set(fmt, db_url, tmp_path):
    out = tmp_path / f"empty.{fmt}"
    count = asyncio.run(
        ex.write_rows(
            ex.fetch_rows(
                "User", url=db_url, table="users", filters={"status": "ghost"}
            ),
            fmt,
            str(out),
            "User export",
        )
    )
    assert count == 0
    assert out.stat().st_size > 0


def test_multi_sheet_workbook_has_one_sheet_per_model(db_url, tmp_path):
    out = tmp_path / "all.xlsx"
    counts = asyncio.run(
        ex.write_xlsx_workbook(
            {
                "User": ex.fetch_rows("User", url=db_url, table="users"),
                "Post": ex.fetch_rows("Post", url=db_url, table="posts"),
            },
            str(out),
        )
    )
    assert counts == {"User": 12, "Post": 4}
    assert_no_leak(out)

    from openpyxl import load_workbook

    assert load_workbook(out).sheetnames == ["User", "Post"]


def test_stream_export_yields_the_whole_file_and_cleans_up(db_url):
    chunks = collect(
        ex.stream_export(
            ex.fetch_rows("User", url=db_url, table="users"), "xlsx", "Users"
        )
    )
    assert chunks and b"".join(chunks)[:2] == b"PK"


def test_serialize_to_temp_then_iter_deletes_the_temp_file(db_url):
    """The generated service's two-step path: fail before streaming, clean up after."""

    async def run() -> tuple[str, bytes]:
        path = await ex.serialize_to_temp(
            ex.fetch_rows("User", url=db_url, table="users"), "xlsx", "Users"
        )
        assert Path(path).exists()
        data = b"".join([c async for c in ex.iter_temp_file(path)])
        return path, data

    path, data = asyncio.run(run())
    assert data[:2] == b"PK"
    assert not Path(path).exists()


def test_serialize_to_temp_leaves_no_file_behind_on_failure(db_url, monkeypatch):
    import tempfile

    created: list[str] = []
    real_mkstemp = tempfile.mkstemp

    def spy(*args, **kwargs):
        handle, path = real_mkstemp(*args, **kwargs)
        created.append(path)
        return handle, path

    monkeypatch.setattr(tempfile, "mkstemp", spy)
    with pytest.raises(ex.ExportError):
        asyncio.run(
            ex.serialize_to_temp(
                ex.fetch_rows("Ghost", url=db_url, table="ghosts"), "xlsx", "G"
            )
        )
    assert created and not Path(created[0]).exists()


def test_missing_writer_library_names_the_package(monkeypatch):
    """A missing writer must be reported by package name, so it can be installed."""
    import builtins

    real_import = builtins.__import__

    def blocked(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("no openpyxl")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked)
    with pytest.raises(ex.ExportDependencyError) as err:
        ex._require("xlsx")
    assert err.value.package == "openpyxl"


# ---------------------------------------------------------------------------
# core/license.py — the stub seam
# ---------------------------------------------------------------------------


def test_pro_gate_is_a_no_op_in_phase_7():
    from kaira.core.license import FEATURE_EXPORT_API, is_pro_enabled

    assert is_pro_enabled(FEATURE_EXPORT_API) is True
    assert is_pro_enabled("anything.at.all") is True


# ---------------------------------------------------------------------------
# FEATURE A — kaira export data
# ---------------------------------------------------------------------------


def _export_app():
    from kaira.commands.export_cmd import app

    return app


def test_export_data_writes_into_a_gitignored_exports_dir(project):
    result = runner.invoke(_export_app(), ["data", "User", "--format", "xlsx"])
    assert result.exit_code == 0, result.output

    files = list((project / "exports").glob("user_*.xlsx"))
    assert len(files) == 1
    assert_no_leak(files[0])
    assert "exports/" in (project / ".gitignore").read_text(encoding="utf-8")


def test_export_data_does_not_duplicate_the_gitignore_entry(project):
    runner.invoke(_export_app(), ["data", "User", "--format", "xlsx"])
    runner.invoke(_export_app(), ["data", "User", "--format", "xlsx"])
    text = (project / ".gitignore").read_text(encoding="utf-8")
    assert text.count("exports/") == 1


@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "docx"])
def test_export_data_cli_never_leaks_in_any_format(fmt, project):
    """The CLI half of the sensitive-field-leak regression."""
    result = runner.invoke(_export_app(), ["data", "User", "--format", fmt])
    assert result.exit_code == 0, result.output
    written = list((project / "exports").glob(f"user_*.{fmt}"))
    assert written
    assert_no_leak(written[0])


def test_export_data_respects_limit_and_fields(project):
    result = runner.invoke(
        _export_app(),
        [
            "data",
            "User",
            "--format",
            "xlsx",
            "--limit",
            "4",
            "--fields",
            "username,email",
        ],
    )
    assert result.exit_code == 0, result.output

    from openpyxl import load_workbook

    sheet = load_workbook(next((project / "exports").glob("user_*.xlsx"))).active
    values = list(sheet.values)
    assert values[0] == ("username", "email")
    assert len(values) == 5  # header + 4 rows


def test_export_data_respects_a_filter(project):
    result = runner.invoke(
        _export_app(), ["data", "User", "--format", "xlsx", "--filter", "status:active"]
    )
    assert result.exit_code == 0, result.output
    assert "6" in result.output


def test_export_data_honours_an_explicit_output_path(project):
    target = project / "reports" / "users.xlsx"
    result = runner.invoke(
        _export_app(), ["data", "User", "--format", "xlsx", "--output", str(target)]
    )
    assert result.exit_code == 0, result.output
    assert target.exists()


def test_export_data_all_xlsx_is_one_workbook(project):
    result = runner.invoke(_export_app(), ["data", "--all", "--format", "xlsx"])
    assert result.exit_code == 0, result.output
    books = list((project / "exports").glob("export_*.xlsx"))
    assert len(books) == 1
    assert_no_leak(books[0])

    from openpyxl import load_workbook

    assert load_workbook(books[0]).sheetnames == ["User", "Post"]


def test_export_data_all_pdf_is_one_file_per_model(project):
    result = runner.invoke(_export_app(), ["data", "--all", "--format", "pdf"])
    assert result.exit_code == 0, result.output
    names = sorted(p.name.split("_")[0] for p in (project / "exports").glob("*.pdf"))
    assert names == ["post", "user"]


def test_export_data_rejects_an_unknown_format(project):
    result = runner.invoke(_export_app(), ["data", "User", "--format", "csv"])
    assert result.exit_code != 0
    assert "csv" in result.output


def test_export_data_rejects_an_unknown_model_with_a_suggestion(project):
    result = runner.invoke(_export_app(), ["data", "Ussr", "--format", "xlsx"])
    assert result.exit_code != 0
    assert "User" in result.output


def test_export_data_rejects_a_model_and_all_together(project):
    result = runner.invoke(_export_app(), ["data", "User", "--all", "--format", "xlsx"])
    assert result.exit_code != 0


def test_export_data_requires_a_target(project):
    result = runner.invoke(_export_app(), ["data", "--format", "xlsx"])
    assert result.exit_code != 0


def test_export_data_rejects_a_malformed_filter(project):
    result = runner.invoke(
        _export_app(), ["data", "User", "--format", "xlsx", "--filter", "statusactive"]
    )
    assert result.exit_code != 0


def test_export_data_surfaces_an_unknown_field(project):
    result = runner.invoke(
        _export_app(), ["data", "User", "--format", "xlsx", "--fields", "nope"]
    )
    assert result.exit_code != 0
    assert "nope" in result.output


def test_export_data_requires_a_database_url(project, monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    result = runner.invoke(_export_app(), ["data", "User", "--format", "xlsx"])
    assert result.exit_code != 0
    assert "DATABASE_URL" in result.output


def test_filter_values_may_contain_a_colon():
    from kaira.commands.export_cmd import _parse_key_values

    parsed = _parse_key_values("created_at:2026-08-03T09:00:00", "--filter")
    assert parsed == {"created_at": "2026-08-03T09:00:00"}


# ---------------------------------------------------------------------------
# FEATURE A — production friction
# ---------------------------------------------------------------------------


def test_bulk_export_in_production_requires_the_typed_project_name(
    project, monkeypatch
):
    monkeypatch.setenv("APP_ENV", "production")
    result = runner.invoke(
        _export_app(), ["data", "--all", "--format", "xlsx"], input="wrong\n"
    )
    assert result.exit_code != 0
    assert not list((project / "exports").glob("*.xlsx"))


def test_force_does_not_bypass_the_production_gate(project, monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    result = runner.invoke(
        _export_app(), ["data", "--all", "--format", "xlsx", "--force"], input="\n"
    )
    assert result.exit_code != 0
    assert not list((project / "exports").glob("*.xlsx"))


def test_bulk_export_in_production_proceeds_once_confirmed(project, monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    result = runner.invoke(
        _export_app(), ["data", "--all", "--format", "xlsx"], input=f"{project.name}\n"
    )
    assert result.exit_code == 0, result.output
    assert list((project / "exports").glob("export_*.xlsx"))


def test_single_model_export_in_production_is_not_gated(project, monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    result = runner.invoke(_export_app(), ["data", "User", "--format", "xlsx"])
    assert result.exit_code == 0, result.output


# ---------------------------------------------------------------------------
# FEATURE B — generated API export endpoints
# ---------------------------------------------------------------------------


@pytest.fixture
def scaffolded(project):
    """A project whose User and Post layers were produced by `kaira generate`."""
    from kaira.commands.generate import app as generate_app

    for model, fields in (
        ("User", "username:str, email:str, password:str, api_key:str, status:str"),
        ("Post", "title:str, body:str"),
    ):
        result = runner.invoke(
            generate_app, ["model", model, "--fields", fields, "--tier", "full"]
        )
        assert result.exit_code == 0, result.output

    config_dir = project / "config"
    config_dir.mkdir(exist_ok=True)
    from jinja2 import Environment, FileSystemLoader

    templates = Path(__file__).parent.parent / "kaira" / "templates"
    env = Environment(
        loader=FileSystemLoader(str(templates)),
        keep_trailing_newline=True,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    (config_dir / "settings.py").write_text(
        env.get_template("env_settings.py.j2").render(
            project_name="Test",
            project_slug="test",
            db_type="sqlite",
            auth_type="jwt",
            api_version="v1",
        ),
        encoding="utf-8",
    )
    return project


def test_export_add_generates_a_streaming_guarded_endpoint(scaffolded):
    result = runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    assert result.exit_code == 0, result.output

    router = (scaffolded / "routers" / "user_router.py").read_text(encoding="utf-8")
    assert '@router.get(\n    "/export"' in router
    assert "StreamingResponse" in router
    assert "current_user: dict = Depends(get_current_user)" in router
    assert "@limiter.limit(settings.EXPORT_RATE_LIMIT)" in router
    assert '"X-Kaira-Export-Format": format' in router
    # A raw exception must never reach the client.
    assert 'detail="Export failed. Please try again later."' in router
    assert "{exc}" not in router.split("HTTPException")[1][:400]

    service = (scaffolded / "services" / "user_service.py").read_text(encoding="utf-8")
    assert "async def export(self, fmt: str)" in service
    assert "serialize_to_temp" in service


def test_generated_files_are_valid_python(scaffolded):
    import py_compile

    runner.invoke(_export_app(), ["add", "User", "--format", "all"])
    for rel in (
        "routers/user_router.py",
        "services/user_service.py",
        "core/export.py",
        "core/security.py",
    ):
        py_compile.compile(str(scaffolded / rel), doraise=True)


def test_export_add_copies_the_shared_pipeline_verbatim(scaffolded):
    """The generated app must run the same strip code, not a second copy of it."""
    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    source_dir = Path(__file__).parent.parent / "kaira" / "core"
    for name in ("export.py", "security.py"):
        assert (scaffolded / "core" / name).read_text(encoding="utf-8") == (
            source_dir / name
        ).read_text(encoding="utf-8")


def test_export_add_refuses_an_unguarded_model(scaffolded):
    """The Phase 4 smart-error path — no auth layer, no export endpoint."""
    router = scaffolded / "routers" / "post_router.py"
    router.write_text(
        router.read_text(encoding="utf-8").replace("get_current_user", "get_nothing"),
        encoding="utf-8",
    )
    result = runner.invoke(_export_app(), ["add", "Post", "--format", "xlsx"])
    assert result.exit_code != 0
    assert "auth" in result.output.lower()
    assert "add-guard" in result.output
    assert "KAIRA_EXPORT" not in router.read_text(encoding="utf-8")


def test_export_add_injects_the_rate_limit_setting(scaffolded):
    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    settings = (scaffolded / "config" / "settings.py").read_text(encoding="utf-8")
    assert 'EXPORT_RATE_LIMIT: str = "5/minute"' in settings


def test_export_add_honours_a_rate_limit_override(scaffolded):
    runner.invoke(
        _export_app(), ["add", "User", "--format", "xlsx", "--rate-limit", "2/hour"]
    )
    settings = (scaffolded / "config" / "settings.py").read_text(encoding="utf-8")
    assert 'EXPORT_RATE_LIMIT: str = "2/hour"' in settings


def test_export_add_refuses_to_stack_a_second_endpoint(scaffolded):
    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    result = runner.invoke(_export_app(), ["add", "User", "--format", "pdf"])
    assert result.exit_code != 0
    assert "already" in result.output.lower()


def test_export_add_force_replaces_rather_than_appends(scaffolded):
    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    result = runner.invoke(_export_app(), ["add", "User", "--format", "all", "--force"])
    assert result.exit_code == 0, result.output

    router = (scaffolded / "routers" / "user_router.py").read_text(encoding="utf-8")
    assert router.count("# [KAIRA_EXPORT_ROUTE]") == 1
    assert router.count("async def export_users") == 1
    assert '"xlsx", "pdf", "docx"' in router


def test_export_add_requires_a_service_layer(scaffolded):
    (scaffolded / "services" / "user_service.py").unlink()
    result = runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    assert result.exit_code != 0
    assert "service" in result.output.lower()


def test_export_list_reports_this_endpoints_own_rate_limit(scaffolded):
    """Not the first @limiter.limit in the file — the export route's own."""
    runner.invoke(_export_app(), ["add", "User", "--format", "all"])
    result = runner.invoke(_export_app(), ["list"])
    assert result.exit_code == 0, result.output
    assert "User" in result.output
    assert "5/minute" in result.output
    assert "20/minute" not in result.output
    assert "✅" in result.output


def test_export_list_is_empty_before_anything_is_added(scaffolded):
    result = runner.invoke(_export_app(), ["list"])
    assert result.exit_code == 0
    assert "No export endpoints" in result.output


def test_export_remove_leaves_no_orphaned_imports(scaffolded):
    router = scaffolded / "routers" / "user_router.py"
    service = scaffolded / "services" / "user_service.py"
    before = router.read_text(encoding="utf-8")

    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])
    result = runner.invoke(_export_app(), ["remove", "User"])
    assert result.exit_code == 0, result.output

    after = router.read_text(encoding="utf-8")
    for orphan in (
        "KAIRA_EXPORT",
        "StreamingResponse",
        "EXPORT_MEDIA_TYPES",
        "Literal",
    ):
        assert orphan not in after
    for orphan in (
        "KAIRA_EXPORT",
        "serialize_to_temp",
        "iter_temp_file",
        "AsyncIterator",
    ):
        assert orphan not in service.read_text(encoding="utf-8")

    import py_compile

    py_compile.compile(str(router), doraise=True)
    py_compile.compile(str(service), doraise=True)
    assert after.strip() == before.strip()


def test_export_remove_reports_when_there_is_nothing_to_remove(scaffolded):
    result = runner.invoke(_export_app(), ["remove", "User"])
    assert result.exit_code != 0
    assert "No generated export endpoint" in result.output


def test_generated_service_export_path_produces_a_clean_file(scaffolded, db_url):
    """End-to-end over the code the endpoint actually calls, minus HTTP."""
    runner.invoke(_export_app(), ["add", "User", "--format", "xlsx"])

    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "generated_export", scaffolded / "core" / "export.py"
    )
    assert spec and spec.loader
    # Loaded through the package the copy expects, so its relative import of
    # `.security` resolves against the generated project's core/ — proving the
    # copy is self-sufficient, not quietly falling back on kaira's own module.
    import sys

    sys.path.insert(0, str(scaffolded))
    try:
        generated = importlib.import_module("core.export")
    finally:
        sys.path.remove(str(scaffolded))

    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    async def run() -> bytes:
        engine = create_async_engine(db_url)
        maker = async_sessionmaker(engine, expire_on_commit=False)
        async with maker() as session:
            rows = generated.fetch_rows("User", db=session, table="users")
            path = await generated.serialize_to_temp(rows, "xlsx", "User export")
            data = b"".join([c async for c in generated.iter_temp_file(path)])
        await engine.dispose()
        return data

    data = asyncio.run(run())
    out = scaffolded / "from_api.xlsx"
    out.write_bytes(data)
    assert_no_leak(out)
    assert generated.EXPORT_MEDIA_TYPES["xlsx"].endswith("spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Guide page
# ---------------------------------------------------------------------------


def test_guide_export_shows_both_paths():
    from kaira.commands.guide_cmd import app as guide_app

    result = runner.invoke(guide_app, ["export"])
    assert result.exit_code == 0
    assert "kaira export data" in result.output
    assert "kaira export add" in result.output


def test_guide_index_lists_export():
    from kaira.commands.guide_cmd import app as guide_app

    result = runner.invoke(guide_app, [])
    assert "kaira guide export" in result.output


# ---------------------------------------------------------------------------
# Document-store path — exercised through a Motor-shaped stub
# ---------------------------------------------------------------------------


class _FakeCursor:
    """The slice of Motor's cursor API that ``_stream_document`` touches."""

    def __init__(self, docs: list[dict]) -> None:
        self._docs = docs
        self.batch_sizes: list[int] = []
        self.limits: list[int] = []

    def batch_size(self, size: int) -> "_FakeCursor":
        self.batch_sizes.append(size)
        return self

    def limit(self, count: int) -> "_FakeCursor":
        self.limits.append(count)
        self._docs = self._docs[:count]
        return self

    async def __aiter__(self):
        for doc in self._docs:
            yield doc


class _FakeCollection:
    def __init__(self, docs: list[dict]) -> None:
        self._docs = docs
        self.query: dict | None = None
        self.projection: dict | None = None
        self.cursor: _FakeCursor | None = None

    def find(self, query, projection):
        self.query = query
        self.projection = projection
        docs = [d for d in self._docs if all(d.get(k) == v for k, v in query.items())]
        self.cursor = _FakeCursor([dict(d) for d in docs])
        return self.cursor


def _mongo_docs() -> list[dict]:
    return [
        {
            "_id": f"objectid-{i}",
            "username": f"user{i}",
            "email": f"u{i}@example.com",
            "password": "LEAKPASSWORD",
            "api_key": "LEAKAPIKEY",
            "status": "active" if i % 2 else "inactive",
        }
        for i in range(1, 6)
    ]


def test_document_path_strips_sensitive_keys_and_stringifies_id():
    collection = _FakeCollection(_mongo_docs())
    rows = collect(ex.fetch_rows("User", collection=collection))
    assert len(rows) == 5
    assert "password" not in rows[0] and "api_key" not in rows[0]
    assert rows[0]["_id"] == "objectid-1"


def test_document_path_passes_batch_size_to_the_cursor():
    collection = _FakeCollection(_mongo_docs())
    collect(ex.fetch_rows("User", collection=collection, batch_size=2))
    assert collection.cursor.batch_sizes == [2]


def test_document_path_pushes_limit_down_to_the_cursor():
    collection = _FakeCollection(_mongo_docs())
    rows = collect(ex.fetch_rows("User", collection=collection, limit=2))
    assert collection.cursor.limits == [2]
    assert len(rows) == 2


def test_document_path_applies_equality_filters():
    collection = _FakeCollection(_mongo_docs())
    rows = collect(
        ex.fetch_rows("User", collection=collection, filters={"status": "active"})
    )
    assert collection.query == {"status": "active"}
    assert rows and all(r["status"] == "active" for r in rows)


def test_document_path_projection_cannot_request_a_secret():
    collection = _FakeCollection(_mongo_docs())
    collect(
        ex.fetch_rows("User", collection=collection, fields=["username", "api_key"])
    )
    assert "api_key" not in collection.projection
    assert collection.projection["username"] == 1


def test_document_path_rejects_an_all_sensitive_allowlist():
    collection = _FakeCollection(_mongo_docs())
    with pytest.raises(ex.ExportError):
        collect(ex.fetch_rows("User", collection=collection, fields=["password"]))


def test_document_path_rejects_a_sensitive_filter():
    """Same oracle rule as SQL, enforced without a schema to check against."""
    collection = _FakeCollection(_mongo_docs())
    with pytest.raises(ex.ExportError):
        collect(
            ex.fetch_rows("User", collection=collection, filters={"api_key": "guess"})
        )


def test_document_path_requires_a_url_when_no_collection_is_given():
    with pytest.raises(ex.ExportError):
        collect(ex.fetch_rows("User", is_document_db=True))


@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "docx"])
def test_document_export_never_leaks_in_any_format(fmt, tmp_path):
    """The document-store half of the leak regression."""
    out = tmp_path / f"docs.{fmt}"
    asyncio.run(
        ex.write_rows(
            ex.fetch_rows("User", collection=_FakeCollection(_mongo_docs())),
            fmt,
            str(out),
            "User export",
        )
    )
    assert_no_leak(out)


# ---------------------------------------------------------------------------
# Missing-writer install prompt
# ---------------------------------------------------------------------------


def test_declining_the_install_prompt_aborts_without_installing(monkeypatch):
    from kaira.commands import export_cmd

    monkeypatch.setattr(export_cmd.typer, "confirm", lambda *a, **k: False)
    assert (
        export_cmd._offer_install(ex.ExportDependencyError("pdf", "reportlab")) is False
    )


def test_accepting_the_install_prompt_uses_the_rich_installer(monkeypatch):
    """Never raw pip output — installs go through the Phase 3 installer."""
    from kaira.commands import export_cmd
    from kaira.commands import project as project_cmd

    calls: list[list[str]] = []
    monkeypatch.setattr(export_cmd.typer, "confirm", lambda *a, **k: True)
    monkeypatch.setattr(
        project_cmd, "install_packages", lambda pkgs: (calls.append(pkgs), (1, 0, 0))[1]
    )
    assert (
        export_cmd._offer_install(ex.ExportDependencyError("docx", "python-docx"))
        is True
    )
    assert calls == [["python-docx"]]


def test_a_missing_writer_is_retried_once_after_install(monkeypatch):
    """The retry rebuilds the coroutine — a spent async generator cannot resume."""
    from kaira.commands import export_cmd

    attempts = {"n": 0}

    async def flaky() -> str:
        attempts["n"] += 1
        if attempts["n"] == 1:
            raise ex.ExportDependencyError("xlsx", "openpyxl")
        return "ok"

    monkeypatch.setattr(export_cmd, "_offer_install", lambda _missing: True)
    assert export_cmd._run_export(flaky) == "ok"
    assert attempts["n"] == 2
