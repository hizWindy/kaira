"""
core/export.py

Shared export pipeline: row-fetch → sensitive-field strip → format serialization.
Used by both the ``kaira export`` CLI and generated API export endpoints.

Both trust boundaries run this one module, so a fix to serialization or to the
sensitive-field rule lands on both at once. ``kaira export add`` copies this
file verbatim into ``<project>/core/export.py``; for that copy to work the
module must import nothing from :mod:`kaira` and reach for third-party
packages only inside the function that needs them.

Memory: :func:`fetch_rows` is always batched and never materialises a table.
``xlsx`` is written in openpyxl's write-only mode, so it streams end to end.
``pdf`` and ``docx`` are page-oriented container formats whose libraries need
the assembled document before they can emit a byte — those two hold the
rendered rows for the duration of the write, which is why ``--limit`` exists.
"""

from __future__ import annotations

import json
import os
import tempfile
from collections.abc import AsyncIterator
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import IO, Any, Literal, Union
from uuid import UUID

from .security import (  # noqa: F401 — re-exported as the public strip API
    SENSITIVE_FIELD_MARKERS,
    is_sensitive_field,
    safe_field_names,
    strip_sensitive_fields,
)

ExportFormat = Literal["xlsx", "pdf", "docx"]

#: Every format the pipeline can emit, in the order they are offered on the CLI.
EXPORT_FORMATS: tuple[str, ...] = ("xlsx", "pdf", "docx")

#: Content types for the ``StreamingResponse`` of a generated export endpoint.
EXPORT_MEDIA_TYPES: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

#: PyPI distribution that provides each writer, surfaced in the install prompt.
EXPORT_PACKAGES: dict[str, str] = {
    "xlsx": "openpyxl",
    "pdf": "reportlab",
    "docx": "python-docx",
}

DEFAULT_BATCH_SIZE = 500

# Rows per platypus table. One table per chunk keeps column-width computation
# bounded; ``repeatRows`` re-prints the header on every page of every chunk.
_PDF_ROWS_PER_TABLE = 250

# Long values wrap badly in a fixed-width table cell and push columns off the
# page, so cell text is elided rather than allowed to reflow the layout.
_MAX_CELL_CHARS = 80

_STREAM_CHUNK_BYTES = 64 * 1024

# Sheet titles are an Excel-level constraint, not ours: 31 chars, and none of
# these characters.
_SHEET_TITLE_MAX = 31
_SHEET_TITLE_BANNED = set(r"[]:*?/\\")

Destination = Union[str, "os.PathLike[str]", IO[bytes]]


class ExportError(RuntimeError):
    """An export could not be produced. Message is safe to show a CLI user."""


class ExportDependencyError(ExportError):
    """A format was requested whose writer library is not installed."""

    def __init__(self, fmt: str, package: str) -> None:
        self.format = fmt
        self.package = package
        super().__init__(
            f"Exporting to .{fmt} needs the '{package}' package, which is not installed."
        )


# ---------------------------------------------------------------------------
# Naming
# ---------------------------------------------------------------------------


def default_table_name(model_name: str) -> str:
    """Derive the table/collection name a model is stored under.

    Mirrors ``kaira.core.parser.table_name`` — duplicated here only because the
    copy of this module inside a generated project cannot import Kaira. Callers
    that *can* reach the parser should pass ``table=`` explicitly so this
    fallback is never the thing that decides.

    Args:
        model_name: PascalCase model name.

    Returns:
        snake_case plural name, e.g. ``UserProfile`` → ``user_profiles``.
    """
    out: list[str] = []
    for index, char in enumerate(model_name):
        if char.isupper() and index and not model_name[index - 1].isupper():
            out.append("_")
        out.append(char.lower())
    snake = "".join(out)
    if snake.endswith("y") and not snake.endswith(("ay", "ey", "oy", "uy")):
        return snake[:-1] + "ies"
    if snake.endswith(("s", "sh", "ch", "x", "z")):
        return snake + "es"
    return snake + "s"


# ---------------------------------------------------------------------------
# Value coercion
# ---------------------------------------------------------------------------


def _scalar(value: Any) -> Any:
    """Reduce a database value to something a document writer can hold."""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return f"<{len(bytes(value))} bytes>"
    if isinstance(value, (dict, list, tuple, set)):
        try:
            return json.dumps(value, default=str, ensure_ascii=False)
        except (TypeError, ValueError):
            return str(value)
    return str(value)


def _text(value: Any) -> str:
    """Render a value as elided display text for a PDF/DOCX table cell."""
    text = str(_scalar(value))
    if len(text) > _MAX_CELL_CHARS:
        return text[: _MAX_CELL_CHARS - 1] + "…"
    return text


def _sheet_title(name: str) -> str:
    """Coerce a model name into a title Excel will accept for a worksheet."""
    cleaned = "".join("_" if c in _SHEET_TITLE_BANNED else c for c in name).strip()
    return (cleaned or "Sheet")[:_SHEET_TITLE_MAX]


# ---------------------------------------------------------------------------
# Field resolution
# ---------------------------------------------------------------------------


def resolve_fields(
    available: list[str],
    fields: list[str] | None = None,
) -> list[str]:
    """Decide which columns an export may read.

    Sensitive fields are dropped last and unconditionally, so an explicit
    ``--fields hashed_password`` narrows the request but cannot widen it.

    Args:
        available: Every field the table/collection actually has.
        fields: Optional caller allowlist. ``None`` means all-minus-sensitive.

    Returns:
        Field names in the caller's order when an allowlist was given, else
        the table's own order.

    Raises:
        ExportError: If a requested field does not exist, or if nothing
            exportable remains.
    """
    exportable = safe_field_names(available)
    if fields:
        unknown = [f for f in fields if f not in available]
        if unknown:
            raise ExportError(
                f"Unknown field(s): {', '.join(unknown)}. "
                f"Available: {', '.join(exportable)}"
            )
        chosen = safe_field_names(fields)
    else:
        chosen = exportable

    if not chosen:
        raise ExportError(
            "No exportable fields remain — every requested field is sensitive "
            "and sensitive fields are never exported."
        )
    return chosen


def validate_filters(available: list[str], filters: dict[str, str] | None) -> None:
    """Reject filter keys that are not real, exportable fields.

    The keys become column references in a SQLAlchemy expression, so an
    unrecognised key is caught here rather than being handed to the database.

    Sensitive fields are rejected as filter keys as firmly as they are as
    output columns: a filter is a yes/no oracle, and equality-testing a
    password hash or an API key one guess at a time reads the value out of a
    table that never returns it.

    Raises:
        ExportError: If any filter key is unknown or sensitive.
    """
    if not filters:
        return
    exportable = safe_field_names(available)
    rejected = [k for k in filters if k not in exportable]
    if rejected:
        raise ExportError(
            f"Cannot filter on: {', '.join(rejected)}. "
            f"Available: {', '.join(exportable)}"
        )


# ---------------------------------------------------------------------------
# Row fetching — relational
# ---------------------------------------------------------------------------


def _coerce_filter_value(column: Any, raw: str) -> Any:
    """Cast a ``--filter`` string to the column's Python type where possible.

    ``status:active`` against a text column stays a string; ``age:30`` against
    an integer column becomes ``30``, because a string there matches nothing on
    a strict engine like PostgreSQL.
    """
    try:
        python_type = column.type.python_type
    except (NotImplementedError, AttributeError):
        return raw
    if python_type is bool:
        return str(raw).strip().lower() in {"1", "true", "yes", "y", "on"}
    if python_type in (int, float, Decimal):
        try:
            return python_type(raw)
        except (TypeError, ValueError, ArithmeticError):
            return raw
    return raw


def _order_column(table: Any) -> Any:
    """Pick a stable sort key so LIMIT/OFFSET paging cannot skip or repeat rows.

    Without an ORDER BY, a database is free to return rows in a different order
    for each batch, which silently drops and duplicates rows across the page
    boundary.
    """
    primary = list(table.primary_key.columns)
    if primary:
        return primary[0]
    columns = list(table.columns)
    return columns[0] if columns else None


async def _reflect_table(conn: Any, table: str) -> Any:
    """Load a live ``Table`` from the database's own metadata."""
    from sqlalchemy import MetaData, Table
    from sqlalchemy.exc import SQLAlchemyError

    metadata = MetaData()

    def _load(sync_conn: Any) -> Any:
        return Table(table, metadata, autoload_with=sync_conn)

    try:
        return await conn.run_sync(_load)
    except SQLAlchemyError as exc:  # table missing, no permission, …
        raise ExportError(
            f"Table '{table}' could not be read. "
            "Run your migrations first (kaira migrate upgrade)."
        ) from exc


async def _stream_relational(
    conn: Any,
    table: str,
    *,
    filters: dict[str, str] | None,
    fields: list[str] | None,
    limit: int | None,
    batch_size: int,
) -> AsyncIterator[dict[str, Any]]:
    """Yield rows from a SQL table in LIMIT/OFFSET batches.

    Queries are built from reflected :class:`~sqlalchemy.Table` columns and
    bound parameters — no fragment of the statement is assembled from user
    text, so a filter value cannot become SQL.
    """
    from sqlalchemy import select

    tbl = await _reflect_table(conn, table)
    available = [c.name for c in tbl.columns]
    validate_filters(available, filters)
    columns = resolve_fields(available, fields)

    order = _order_column(tbl)
    remaining = limit
    offset = 0

    while True:
        page = batch_size if remaining is None else min(batch_size, remaining)
        if page <= 0:
            return

        stmt = select(*[tbl.c[name] for name in columns])
        for key, value in (filters or {}).items():
            column = tbl.c[key]
            stmt = stmt.where(column == _coerce_filter_value(column, value))
        if order is not None:
            stmt = stmt.order_by(order)
        stmt = stmt.limit(page).offset(offset)

        result = await conn.execute(stmt)
        batch = result.mappings().all()
        if not batch:
            return

        for row in batch:
            # Stripped again on the way out even though the SELECT already
            # excluded sensitive columns: two independent gates, so a future
            # change to column resolution cannot leak by itself.
            yield strip_sensitive_fields(dict(row))

        if remaining is not None:
            remaining -= len(batch)
            if remaining <= 0:
                return
        if len(batch) < page:
            return
        offset += len(batch)


# ---------------------------------------------------------------------------
# Row fetching — document stores
# ---------------------------------------------------------------------------


async def _stream_document(
    collection: Any,
    *,
    filters: dict[str, str] | None,
    fields: list[str] | None,
    limit: int | None,
    batch_size: int,
) -> AsyncIterator[dict[str, Any]]:
    """Yield documents from a Motor collection using server-side cursor batching.

    A document store has no schema to validate a filter key against, so the
    only check that can be made here is the one that matters: a sensitive key
    is never accepted as a filter, for the same oracle reason as SQL.
    """
    query: dict[str, Any] = dict(filters or {})
    sensitive_keys = [k for k in query if is_sensitive_field(k)]
    if sensitive_keys:
        raise ExportError(f"Cannot filter on: {', '.join(sensitive_keys)}.")

    projection: dict[str, int] | None = None
    if fields:
        allowed = safe_field_names(fields)
        if not allowed:
            raise ExportError(
                "No exportable fields remain — every requested field is sensitive "
                "and sensitive fields are never exported."
            )
        projection = {name: 1 for name in allowed}
        projection.setdefault("_id", 0)

    cursor = collection.find(query, projection)
    cursor = cursor.batch_size(batch_size)
    if limit:
        cursor = cursor.limit(limit)

    async for doc in cursor:
        row = dict(doc)
        if "_id" in row:
            row["_id"] = str(row["_id"])
        yield strip_sensitive_fields(row)


# ---------------------------------------------------------------------------
# Row fetching — public entry point
# ---------------------------------------------------------------------------


async def fetch_rows(
    model_name: str,
    *,
    filters: dict[str, str] | None = None,
    fields: list[str] | None = None,
    limit: int | None = None,
    batch_size: int = DEFAULT_BATCH_SIZE,
    db: Any = None,
    collection: Any = None,
    url: str = "",
    table: str = "",
    is_document_db: bool = False,
) -> AsyncIterator[dict[str, Any]]:
    """Stream rows for a model in batches. Never loads the full table into memory.

    Args:
        model_name: PascalCase model name as registered in ``.kaira.json``.
        filters: equality filters, e.g. ``{"status": "active"}``
        fields: explicit field allowlist; None means all-minus-sensitive
        limit: optional row cap
        batch_size: rows fetched per DB round-trip
        db: An open ``AsyncSession``/``AsyncConnection``. Passed by a generated
            endpoint so the export rides the app's existing pool instead of
            opening a second one per request.
        collection: An open Motor collection — the document-store equivalent of
            ``db`` (``User.get_motor_collection()`` inside a generated app).
        url: Connection URL, used by the CLI where no session exists.
        table: Table/collection name. Defaults to :func:`default_table_name`.
        is_document_db: Select the Motor path instead of the SQL path.

    Yields:
        One row dict at a time, sensitive fields already stripped.

    Raises:
        ExportError: On an unreadable table, an unknown field/filter, or when
            neither a live handle nor a URL was supplied.
    """
    target = table or default_table_name(model_name)

    if collection is not None:
        async for row in _stream_document(
            collection,
            filters=filters,
            fields=fields,
            limit=limit,
            batch_size=batch_size,
        ):
            yield row
        return

    if is_document_db:
        if not url:
            raise ExportError("No database URL configured — cannot read documents.")
        from motor.motor_asyncio import AsyncIOMotorClient

        client: Any = AsyncIOMotorClient(url, serverSelectionTimeoutMS=5000)
        try:
            database = client.get_default_database()
            async for row in _stream_document(
                database[target],
                filters=filters,
                fields=fields,
                limit=limit,
                batch_size=batch_size,
            ):
                yield row
        finally:
            client.close()
        return

    if db is not None:
        # An AsyncSession hands out the AsyncConnection it is already bound to;
        # an AsyncConnection is one already.
        conn = await db.connection() if hasattr(db, "connection") else db
        async for row in _stream_relational(
            conn,
            target,
            filters=filters,
            fields=fields,
            limit=limit,
            batch_size=batch_size,
        ):
            yield row
        return

    if not url:
        raise ExportError("No database URL configured — cannot read rows.")

    from sqlalchemy.ext.asyncio import create_async_engine

    engine = create_async_engine(url)
    try:
        async with engine.connect() as conn:
            async for row in _stream_relational(
                conn,
                target,
                filters=filters,
                fields=fields,
                limit=limit,
                batch_size=batch_size,
            ):
                yield row
    finally:
        await engine.dispose()


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def _require(fmt: str) -> None:
    """Fail early, and by name, when a writer library is missing."""
    package = EXPORT_PACKAGES[fmt]
    module = {"xlsx": "openpyxl", "pdf": "reportlab", "docx": "docx"}[fmt]
    try:
        __import__(module)
    except ImportError as exc:
        raise ExportDependencyError(fmt, package) from exc


async def _collect(
    rows: AsyncIterator[dict[str, Any]],
) -> tuple[list[str], list[list[Any]]]:
    """Drain an async row iterator into a header and a list of value rows.

    Used by the writers whose libraries cannot accept rows incrementally.
    The header comes from the first row: a table has a fixed shape, and a
    document store's first document defines the export's shape so later
    documents cannot silently widen the file mid-write.
    """
    header: list[str] = []
    values: list[list[Any]] = []
    async for row in rows:
        if not header:
            header = list(row.keys())
        values.append([row.get(name) for name in header])
    return header, values


async def write_xlsx(
    rows: AsyncIterator[dict[str, Any]],
    output_path: Destination,
    title: str = "Export",
) -> int:
    """Stream rows into an .xlsx file using openpyxl write-only mode.

    Args:
        rows: Async iterator of already-stripped row dicts.
        output_path: Filesystem path or open binary file object.
        title: Worksheet name.

    Returns:
        Number of data rows written.
    """
    _require("xlsx")
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=_sheet_title(title))
    count = await _append_sheet(sheet, rows)
    workbook.save(output_path)
    return count


async def _append_sheet(sheet: Any, rows: AsyncIterator[dict[str, Any]]) -> int:
    """Append a header plus every row to a write-only worksheet."""
    header: list[str] = []
    count = 0
    async for row in rows:
        if not header:
            header = list(row.keys())
            sheet.append(header)
        sheet.append([_scalar(row.get(name)) for name in header])
        count += 1
    if not header:
        sheet.append(["(no rows)"])
    return count


async def write_xlsx_workbook(
    sheets: dict[str, AsyncIterator[dict[str, Any]]],
    output_path: Destination,
) -> dict[str, int]:
    """Write one workbook with one sheet per model — the ``--all`` xlsx shape.

    Args:
        sheets: Mapping of model name → async row iterator.
        output_path: Filesystem path or open binary file object.

    Returns:
        Mapping of model name → rows written.
    """
    _require("xlsx")
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    counts: dict[str, int] = {}
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=_sheet_title(name))
        counts[name] = await _append_sheet(sheet, rows)
    if not counts:
        workbook.create_sheet(title="Empty").append(["(no models)"])
    workbook.save(output_path)
    return counts


async def write_pdf(
    rows: AsyncIterator[dict[str, Any]],
    output_path: Destination,
    title: str,
) -> int:
    """Render rows into a paginated table using reportlab platypus.

    Args:
        rows: Async iterator of already-stripped row dicts.
        output_path: Filesystem path or open binary file object.
        title: Document heading.

    Returns:
        Number of data rows written.
    """
    _require("pdf")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    header, values = await _collect(rows)
    styles = getSampleStyleSheet()

    document = SimpleDocTemplate(
        output_path if isinstance(output_path, (str, os.PathLike)) else output_path,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )

    story: list[Any] = [
        Paragraph(title, styles["Title"]),
        Paragraph(f"{len(values)} record(s) · generated by Kaira", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    if not header:
        story.append(Paragraph("No records matched this export.", styles["Normal"]))
    else:
        table_style = TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9ca3af")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f3f4f6")],
                ),
            ]
        )
        head = [_text(name) for name in header]
        for start in range(0, max(len(values), 1), _PDF_ROWS_PER_TABLE):
            chunk = values[start : start + _PDF_ROWS_PER_TABLE]
            data = [head] + [[_text(cell) for cell in row] for row in chunk]
            block = Table(data, repeatRows=1)
            block.setStyle(table_style)
            if start:
                story.append(PageBreak())
            story.append(block)

    document.build(story)
    return len(values)


async def write_docx(
    rows: AsyncIterator[dict[str, Any]],
    output_path: Destination,
    title: str,
) -> int:
    """Render rows into a table using python-docx.

    Args:
        rows: Async iterator of already-stripped row dicts.
        output_path: Filesystem path or open binary file object.
        title: Document heading.

    Returns:
        Number of data rows written.
    """
    _require("docx")
    from docx import Document

    header, values = await _collect(rows)
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"{len(values)} record(s) · generated by Kaira")

    if not header:
        document.add_paragraph("No records matched this export.")
    else:
        table = document.add_table(rows=1, cols=len(header))
        # A built-in style name that ships with the default python-docx
        # template; anything custom would not exist in the blank document.
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:  # pragma: no cover — template without that style
            pass
        for cell, name in zip(table.rows[0].cells, header):
            cell.text = _text(name)
        for row in values:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = _text(value)

    # python-docx takes a path *string* or a file object, not a PathLike.
    document.save(
        str(output_path) if isinstance(output_path, os.PathLike) else output_path
    )
    return len(values)


async def write_rows(
    rows: AsyncIterator[dict[str, Any]],
    fmt: str,
    output_path: Destination,
    title: str,
) -> int:
    """Serialize rows in *fmt* to *output_path*.

    The one place a format string turns into a writer, so the CLI and the
    generated endpoint cannot end up supporting different sets.

    Raises:
        ExportError: If *fmt* is not a supported format.
    """
    if fmt == "xlsx":
        return await write_xlsx(rows, output_path, title=title)
    if fmt == "pdf":
        return await write_pdf(rows, output_path, title)
    if fmt == "docx":
        return await write_docx(rows, output_path, title)
    raise ExportError(
        f"Unsupported export format '{fmt}'. Supported: {', '.join(EXPORT_FORMATS)}"
    )


async def serialize_to_temp(
    rows: AsyncIterator[dict[str, Any]],
    fmt: str,
    title: str,
) -> str:
    """Serialize rows to a temporary file and return its path.

    Deliberately not a generator. A generated endpoint calls this *before*
    returning its ``StreamingResponse``, so a failure — an unreadable table, a
    missing writer library — still happens while the status code can be
    changed. Raise from inside the streaming iterator instead and the client
    has already been told ``200 OK``, then receives a truncated file.

    The caller owns the returned path and must delete it;
    :func:`iter_temp_file` does that as it finishes reading.
    """
    handle, temp_path = tempfile.mkstemp(prefix="kaira_export_", suffix=f".{fmt}")
    os.close(handle)
    try:
        await write_rows(rows, fmt, temp_path, title)
    except BaseException:
        _unlink(temp_path)
        raise
    return temp_path


async def iter_temp_file(
    path: str,
    chunk_size: int = _STREAM_CHUNK_BYTES,
) -> AsyncIterator[bytes]:
    """Yield a file's bytes in chunks, deleting it when the stream ends.

    Cleanup sits in a ``finally`` so a client that disconnects mid-download
    does not leave the file behind.
    """
    try:
        with open(path, "rb") as fh:
            while True:
                chunk = fh.read(chunk_size)
                if not chunk:
                    break
                yield chunk
    finally:
        _unlink(path)


def _unlink(path: str) -> None:
    """Delete a temp file, ignoring a failure to do so."""
    try:
        Path(path).unlink()
    except OSError:  # pragma: no cover — best effort cleanup
        pass


async def stream_export(
    rows: AsyncIterator[dict[str, Any]],
    fmt: str,
    title: str,
    chunk_size: int = _STREAM_CHUNK_BYTES,
) -> AsyncIterator[bytes]:
    """Serialize rows and yield the resulting file in chunks.

    The one-call form of :func:`serialize_to_temp` + :func:`iter_temp_file`,
    for callers that are not an HTTP handler and so have no status code to
    protect. The document is assembled on disk rather than in a ``BytesIO`` so
    the finished bytes never sit in the process heap.

    Args:
        rows: Async iterator of already-stripped row dicts.
        fmt: One of :data:`EXPORT_FORMATS`.
        title: Document/sheet heading.
        chunk_size: Bytes per yielded chunk.

    Yields:
        Successive byte chunks of the serialized document.
    """
    temp_path = await serialize_to_temp(rows, fmt, title)
    async for chunk in iter_temp_file(temp_path, chunk_size):
        yield chunk
